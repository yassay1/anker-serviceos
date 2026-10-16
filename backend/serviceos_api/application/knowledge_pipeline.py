"""Upload-driven knowledge pipeline: convert -> chunk -> embed, lease-driven.

Concept-ported from the reference project's ingestion path (document/version
state machine, content-hash dedupe, row-never-split tables, parent chunks) onto
the existing ServiceOS worker lease pattern: no Redis, no external scheduler,
no Spring events. Original bytes live in the evidence object store; the database
holds metadata, segments and vectors only.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import PurePosixPath
from uuid import UUID

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from serviceos_api.application.business import BusinessService
from serviceos_api.application.contracts import (
    KnowledgeDocumentDetail, KnowledgeDocumentRead, KnowledgePage, KnowledgeSearchRead,
    KnowledgeSegmentCounts, KnowledgeStatsRead,
)
from serviceos_api.application.errors import ServiceError
from serviceos_api.application.evidence_storage import EvidenceStorage, EvidenceStorageFailure
from serviceos_api.application.knowledge_retrieval import KnowledgeRetrieval
from serviceos_api.application.knowledge_splitter import Chunk, split_length, split_markdown_header, split_rows
from serviceos_api.integrations.embedding_adapter import EmbeddingAdapter, EmbeddingUnavailable
from serviceos_api.persistence.knowledge_models import (
    EMBEDDING_DIMENSIONS, KnowledgeDocument, KnowledgeDocumentVersion, KnowledgeEmbedding,
    KnowledgeJob, KnowledgeSegment,
)
from serviceos_api.persistence.models import OutboxEvent
from serviceos_api.settings import Settings

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
UPLOAD_TYPES = {
    ".md": ("text/markdown", "MARKDOWN_HEADER"),
    ".markdown": ("text/markdown", "MARKDOWN_HEADER"),
    ".txt": ("text/plain", "MARKDOWN_HEADER"),
    ".xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "ROW"),
    ".csv": ("text/csv", "ROW"),
    ".jsonl": ("application/jsonl", "ROW"),
    ".jpg": ("image/jpeg", "IMAGE"),
    ".jpeg": ("image/jpeg", "IMAGE"),
    ".png": ("image/png", "IMAGE"),
    ".webp": ("image/webp", "IMAGE"),
}
UPLOAD_CURATION_LABEL = ""  # Retired: uploads enter retrieval on success, source_ref keeps provenance.
EMBED_BATCH_SIZE = 10

ORDER_COLUMNS = {
    "订单号": "order_no", "sku": "sku", "product description": "description", "qty": "qty",
    "购买日期": "purchase_date", "质保结束日期": "warranty_end", "订单地址": "address",
    "订单状态": "status", "物流承运商": "carrier", "物流单号": "tracking",
    "物流状态": "logistics_status", "发货日期": "ship_date", "订单金额": "amount",
}


def normalize_seller(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def detect_sheet_semantic(header: list[str]) -> str:
    lowered = {str(h).strip().casefold() for h in header if h}
    if lowered & {"country"} and lowered & {"seller"}:
        return "DEALER"
    if "订单号" in lowered:
        return "ORDER"
    if {"title", "content"} <= lowered:
        return "QA"
    return "ROW"


def _facet_value(key: str, raw) -> str:
    if raw is None:
        return ""
    if hasattr(raw, "strftime"):
        return raw.date().isoformat()
    return " ".join(str(raw).split())


def _order_attrs(header: list[str], row) -> dict:
    attrs = {}
    for cell, column in zip(row, header):
        target = ORDER_COLUMNS.get(str(column).strip().casefold())
        if not target:
            continue
        value = _facet_value(target, cell)
        if value:
            attrs[target] = value.casefold() if target == "order_no" else value
    return attrs


def _dealer_attrs(header: list[str], row) -> dict:
    pairs = {str(h).strip().casefold(): v for h, v in zip(header, row) if h}
    country = str(pairs.get("country") or "").strip().upper()
    seller = str(pairs.get("seller") or "").strip()
    if not country or not seller:
        return {}
    return {"country": country, "seller": seller, "seller_norm": normalize_seller(seller)}


def _embed_content(storage: EvidenceStorage, document: KnowledgeDocument,
                   segment: KnowledgeSegment) -> dict:
    if segment.image_key:
        payload = storage.read(segment.image_key, MAX_UPLOAD_BYTES)
        encoded = base64.b64encode(payload).decode("ascii")
        return {"image": f"data:{document.media_type};base64,{encoded}"}
    return {"text": f"{segment.title}\n{segment.content}"}


def storage_key_for(document_id: UUID, version_id: UUID) -> str:
    return f"{document_id}/{version_id}"


def _event(session: Session, tenant_id: UUID | None, document: KnowledgeDocument) -> None:
    session.add(OutboxEvent(
        tenant_id=tenant_id, conversation_id=None, type="knowledge.document_updated",
        visibility="INTERNAL", data={
            "document_id": str(document.id), "status": document.status,
            "version": document.current_version,
            "error_code": document.error_code,
        },
    ))


class KnowledgePipelineService(BusinessService):
    def __init__(self, session: Session, actor, storage: EvidenceStorage):
        super().__init__(session, actor)
        self.storage = storage

    def _staff(self) -> None:
        if self.actor.kind != "STAFF":
            raise ServiceError(403, "仅客服身份可以管理知识文档。", code="FORBIDDEN_ACTOR")

    def upload(self, filename: str, payload: bytes, title: str | None,
               sheet_overrides: dict | None = None) -> tuple[KnowledgeDocumentRead, bool]:
        self._staff()
        suffix = PurePosixPath(filename or "").suffix.lower()
        if suffix not in UPLOAD_TYPES:
            raise ServiceError(415, "仅支持 md/txt/xlsx/csv/jsonl/jpeg/png/webp 知识文件。",
                               code="UNSUPPORTED_KNOWLEDGE_TYPE")
        if not payload or len(payload) > MAX_UPLOAD_BYTES:
            raise ServiceError(413, "知识文件需在 1 字节到 10MB 之间。", code="KNOWLEDGE_SIZE")
        overrides = {}
        for sheet_name, semantic in (sheet_overrides or {}).items():
            if semantic not in ("ORDER", "DEALER", "QA", "ROW"):
                raise ServiceError(422, "sheet 语义覆盖值不合法。", code="INVALID_SHEET_OVERRIDE")
            overrides[str(sheet_name)] = semantic
        media_type, split_type = UPLOAD_TYPES[suffix]
        content_hash = hashlib.sha256(payload).hexdigest()
        existing = self.session.scalar(select(KnowledgeDocument).where(
            KnowledgeDocument.content_hash == content_hash))
        if existing is not None:
            return self._read(existing), True
        document = KnowledgeDocument(
            source_kind="upload", title=(title or filename).strip()[:200] or filename,
            filename=filename, media_type=media_type, content_hash=content_hash,
            status="UPLOADED", current_version=1, split_type=split_type,
            sheet_overrides=overrides,
            uploaded_by_staff_id=self.actor.actor_id,
        )
        self.session.add(document)
        self.session.flush()
        version = KnowledgeDocumentVersion(document_id=document.id, version=1,
                                           content_hash=content_hash, status="UPLOADED")
        self.session.add(version)
        self.session.flush()
        key = storage_key_for(document.id, version.id)
        try:
            self.storage.put(key, payload)
        except EvidenceStorageFailure as error:
            self.session.rollback()
            raise ServiceError(503, "知识存储暂不可用。", code="KNOWLEDGE_STORAGE_UNAVAILABLE") from error
        try:
            self.session.add(KnowledgeJob(document_id=document.id, version_id=version.id,
                                          tenant_id=self.actor.tenant_id, status="QUEUED"))
            _event(self.session, self.actor.tenant_id, document)
            self.session.commit()
        except Exception:
            self.session.rollback()
            try:
                self.storage.delete(key)
            except EvidenceStorageFailure:
                pass
            raise
        return self._read(document), False

    def reprocess(self, document_id: UUID) -> KnowledgeDocumentRead:
        self._staff()
        if self.actor.role not in ("SUPERVISOR", "ADMIN"):
            raise ServiceError(403, "仅主管或管理员可以重新入队知识文档。", code="FORBIDDEN_ROLE")
        document = self.session.scalar(select(KnowledgeDocument).where(
            KnowledgeDocument.id == document_id).with_for_update())
        if document is None:
            raise ServiceError(404, "知识文档不存在。", code="KNOWLEDGE_MISSING")
        if document.status != "FAILED":
            raise ServiceError(409, "仅失败文档可以重新入队。", code="KNOWLEDGE_NOT_FAILED")
        job = self.session.scalar(select(KnowledgeJob).where(
            KnowledgeJob.document_id == document.id,
            KnowledgeJob.version_id.in_(select(KnowledgeDocumentVersion.id).where(
                KnowledgeDocumentVersion.document_id == document.id,
                KnowledgeDocumentVersion.version == document.current_version)),
        ).with_for_update())
        if job is None or job.status not in ("FAILED", "RUNNING"):
            raise ServiceError(409, "文档作业状态不允许重新入队。", code="KNOWLEDGE_JOB_STATE")
        version = self.session.scalar(select(KnowledgeDocumentVersion).where(
            KnowledgeDocumentVersion.id == job.version_id).with_for_update())
        for segment in self.session.scalars(select(KnowledgeSegment).where(
                KnowledgeSegment.version_id == version.id)).all():
            self.session.delete(segment)
        now = datetime.now(timezone.utc)
        job.status, job.attempt_count, job.available_at = "QUEUED", 0, now
        job.lease_owner = job.lease_token = job.lease_expires_at = None
        job.last_error_code = None
        document.status, document.error_code = "UPLOADED", None
        version.status, version.error_code = "UPLOADED", None
        _event(self.session, job.tenant_id, document)
        self.session.commit()
        return self._read(document)

    def list_documents(self, cursor: UUID | None, limit: int) -> KnowledgePage:
        self._staff()
        query = select(KnowledgeDocument)
        if cursor is not None:
            current = self.session.scalar(select(KnowledgeDocument).where(KnowledgeDocument.id == cursor))
            if current is None:
                raise ServiceError(404, "知识文档不存在。", code="KNOWLEDGE_MISSING")
            query = query.where(KnowledgeDocument.created_at > current.created_at)
        rows = self.session.scalars(query.order_by(KnowledgeDocument.created_at, KnowledgeDocument.id)
                                    .limit(limit + 1)).all()
        return KnowledgePage(items=[self._read(row) for row in rows[:limit]],
                             next_cursor=str(rows[limit].id) if len(rows) > limit else None)

    def get_document(self, document_id: UUID) -> KnowledgeDocumentDetail:
        self._staff()
        document = self.session.scalar(select(KnowledgeDocument).where(KnowledgeDocument.id == document_id))
        if document is None:
            raise ServiceError(404, "知识文档不存在。", code="KNOWLEDGE_MISSING")
        counts = self.session.execute(select(
            func.count(KnowledgeSegment.id),
            func.count(KnowledgeEmbedding.segment_id),
            func.count(KnowledgeSegment.id).filter(KnowledgeSegment.skip_embedding),
        ).select_from(KnowledgeSegment).outerjoin(
            KnowledgeEmbedding, KnowledgeEmbedding.segment_id == KnowledgeSegment.id,
        ).where(KnowledgeSegment.document_id == document.id)).one()
        return KnowledgeDocumentDetail(
            **self._read(document).model_dump(),
            counts=KnowledgeSegmentCounts(segments=counts[0], embedded=counts[1], parents=counts[2]),
        )

    def search(self, query: str, limit: int):
        self._staff()
        retrieval = KnowledgeRetrieval(self.session, Settings.from_environment(), top_results=limit)
        result = retrieval.search(query, set(), gate=False)
        return [KnowledgeSearchRead(entry_key=item.entry_key, source_ref=item.source_ref,
                                    applicability=item.applicability, path=item.path,
                                    excerpt=item.excerpt) for item in result.items]

    def stats(self) -> KnowledgeStatsRead:
        self._staff()
        documents = self.session.scalar(select(func.count()).select_from(KnowledgeDocument)) or 0
        segments = self.session.scalar(select(func.count()).select_from(KnowledgeSegment)) or 0
        embedded = self.session.scalar(select(func.count()).select_from(KnowledgeEmbedding)) or 0
        settings = Settings.from_environment()
        return KnowledgeStatsRead(
            corpus_present=bool(segments), documents=documents, segments=segments,
            embedded=embedded, vector_enabled=settings.embedding_configured,
            provider=settings.embedding_model if settings.embedding_configured else "")

    @staticmethod
    def _read(document: KnowledgeDocument) -> KnowledgeDocumentRead:
        return KnowledgeDocumentRead(
            document_id=document.id, title=document.title, filename=document.filename,
            source_kind=document.source_kind, status=document.status, error_code=document.error_code,
            current_version=document.current_version, split_type=document.split_type,
            uploaded_by_staff_id=document.uploaded_by_staff_id, created_at=document.created_at,
        )


def _sheet_seeds(sheet_label: str, header: list, rows: list, semantic: str) -> list[dict]:
    seeds = []
    clean_header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
    for index, row in enumerate(rows, start=2):
        pairs = [f"{clean_header[i]}: {' '.join(str(cell).split())}"
                 for i, cell in enumerate(row)
                 if i < len(clean_header) and cell is not None and str(cell).strip()]
        if not pairs:
            continue
        attrs: dict = {}
        if semantic == "ORDER":
            attrs = _order_attrs(clean_header, row)
        elif semantic == "DEALER":
            attrs = _dealer_attrs(clean_header, row)
        if semantic == "DEALER" and not attrs:
            continue
        seeds.append({"title": f"{sheet_label} row {index}", "content": " | ".join(pairs),
                      "attrs": attrs})
    return seeds


def _convert_workbook(payload: bytes, overrides: dict) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    seeds = []
    for sheet_name in workbook.sheetnames:
        rows = [row for row in workbook[sheet_name].iter_rows(values_only=True)
                if any(cell is not None and str(cell).strip() for cell in row)]
        if not rows:
            continue
        semantic = overrides.get(sheet_name) or detect_sheet_semantic([str(h) for h in rows[0]])
        seeds.extend(_sheet_seeds(sheet_name, list(rows[0]), rows[1:], semantic))
    return seeds


def _convert_csv(payload: bytes, overrides: dict) -> list[dict]:
    text = payload.decode("utf-8-sig")
    rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    if not rows:
        return []
    semantic = overrides.get("default") or detect_sheet_semantic(rows[0])
    return _sheet_seeds("csv", rows[0], rows[1:], semantic)


def _convert_jsonl(payload: bytes) -> list[dict]:
    seeds = []
    for line_number, line in enumerate(payload.decode("utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except ValueError:
            raise ValueError("jsonl_parse_error") from None
        if not isinstance(row, dict):
            raise ValueError("jsonl_row_not_object")
        title = str(row.get("title") or f"row {line_number}")
        content = str(row.get("content") or "")
        if content.strip():
            seeds.append({"title": title, "content": " ".join(content.split()), "attrs": {}})
    return seeds


def _segment_from_chunk(chunk: Chunk, document: KnowledgeDocument, version: KnowledgeDocumentVersion,
                      index: int, parent_id: UUID | None) -> KnowledgeSegment:
    return KnowledgeSegment(
        document_id=document.id, version_id=version.id, chunk_index=index,
        parent_segment_id=parent_id, skip_embedding=chunk.parent,
        entry_key=f"doc:{document.id}:v{version.version}:{'p' if chunk.parent else 'c'}{index}",
        title=chunk.title, content=chunk.content,
        content_sha256=hashlib.sha256(f"{chunk.title}\n{chunk.content}".encode("utf-8")).hexdigest(),
        models=[], directories=[], pdt=None, source_status="UPLOADED",
        consumer_visible="AGENT_INTERNAL", applicability="STAFF_UPLOAD",
        curation={"source_kind": "upload", "source_ref": document.filename or document.title},
        attrs=chunk.attrs,
        image_key=storage_key_for(document.id, version.id) if chunk.image else None,
    )


def process_knowledge_job(session: Session, job: KnowledgeJob, settings: Settings,
                          storage: EvidenceStorage, embedder: EmbeddingAdapter | None = None) -> str | None:
    """Advance one pipeline job; returns a stable error code or None on success."""
    document = session.scalar(select(KnowledgeDocument).where(
        KnowledgeDocument.id == job.document_id).with_for_update())
    version = session.scalar(select(KnowledgeDocumentVersion).where(
        KnowledgeDocumentVersion.id == job.version_id).with_for_update())
    if document is None or version is None:
        return "knowledge_subject_missing"
    if document.status in ("EMBEDDED", "STORED"):
        return None
    key = storage_key_for(document.id, version.id)
    try:
        payload = storage.read(key, MAX_UPLOAD_BYTES)
    except EvidenceStorageFailure:
        return "knowledge_storage_missing"
    if len(payload) > MAX_UPLOAD_BYTES:
        return "knowledge_object_too_large"
    try:
        overrides = dict(document.sheet_overrides or {})
        seeds: list[dict] | None
        if document.split_type == "IMAGE":
            seeds = [{"title": document.title, "content": document.title, "attrs": {}, "image": True}]
        elif document.split_type == "ROW" and document.media_type == "text/csv":
            seeds = _convert_csv(payload, overrides)
        elif document.split_type == "ROW" and document.media_type.endswith("spreadsheetml.sheet"):
            seeds = _convert_workbook(payload, overrides)
        elif document.split_type == "ROW":
            seeds = _convert_jsonl(payload)
        else:
            seeds = None
            text = payload.decode("utf-8")
        document.status = version.status = "CONVERTED"
        document.error_code = version.error_code = None
        session.flush()

        for segment in session.scalars(select(KnowledgeSegment).where(
                KnowledgeSegment.version_id == version.id)).all():
            session.delete(segment)
        session.flush()
        if seeds is not None:
            chunks = [Chunk(title=seed["title"], content=seed["content"],
                            attrs=seed.get("attrs", {}), image=bool(seed.get("image")))
                      for seed in seeds]
        elif document.split_type == "MARKDOWN_HEADER":
            chunks = split_markdown_header(text, document.title)
        else:
            chunks = split_length(text, document.title)
        if not chunks:
            raise ValueError("no_segments")
        children: list[KnowledgeSegment] = []
        current_parent: UUID | None = None
        for index, chunk in enumerate(chunks):
            segment = _segment_from_chunk(chunk, document, version, index,
                                          None if chunk.parent else current_parent)
            session.add(segment)
            session.flush()
            if chunk.parent:
                current_parent = segment.id
            else:
                children.append(segment)
        document.status = version.status = "CHUNKED"
        session.flush()

        embedder = embedder or EmbeddingAdapter(settings)
        final = "STORED"
        if embedder.configured:
            embeddable = children if embedder.supports_image else [
                segment for segment in children if segment.image_key is None]
            for start in range(0, len(embeddable), EMBED_BATCH_SIZE):
                batch = embeddable[start:start + EMBED_BATCH_SIZE]
                vectors = embedder.embed_contents([_embed_content(storage, document, segment)
                                                   for segment in batch])
                for segment, vector in zip(batch, vectors):
                    session.add(KnowledgeEmbedding(
                        segment_id=segment.id, model=settings.embedding_model,
                        dimensions=EMBEDDING_DIMENSIONS,
                        embedding="[" + ",".join(repr(float(value)) for value in vector) + "]",
                    ))
            final = "EMBEDDED" if len(embeddable) == len(children) else "STORED"
        document.status = version.status = final
        _event(session, job.tenant_id, document)
        session.flush()
        return None
    except (EmbeddingUnavailable, RuntimeError):
        code = "knowledge_embed_failed"
    except Exception:
        code = "knowledge_process_failed"
    document.status = version.status = "FAILED"
    document.error_code = version.error_code = code
    _event(session, job.tenant_id, document)
    session.flush()
    return code
