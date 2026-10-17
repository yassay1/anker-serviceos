"""Structured facets, typed tools, sheet routing and multimodal image ingestion."""

import base64
import io
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from serviceos_api.application.agent_services import AgentToolService
from serviceos_api.application.auth import Principal
from serviceos_api.application.knowledge_pipeline import (
    KnowledgePipelineService, process_knowledge_job, shipment_attrs,
)
from serviceos_api.application.knowledge_retrieval import KnowledgeRetrieval
from serviceos_api.application.tool_envelope import TOOL_STATUSES
from serviceos_api.bootstrap_local import seed_local_accounts
from serviceos_api.integrations.embedding_adapter import EmbeddingAdapter
from serviceos_api.integrations.local_evidence_storage import LocalEvidenceStorage
from serviceos_api.persistence.knowledge_models import (
    KnowledgeDocument, KnowledgeEmbedding, KnowledgeJob, KnowledgeSegment,
)
from serviceos_api.persistence.models import AgentRun, Customer, StaffUser, Tenant
from serviceos_api.settings import Settings
from test_knowledge_retrieval import UNIT_VECTOR, active_run

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def test_split_shipment_projection_does_not_guess_when_source_counts_disagree():
    rows = shipment_attrs({
        "tracking": "Tracking number1: T1 Tracking number2: T2",
        "carrier": "Carrier A Carrier B",
        "ship_date": "2026-09-17",
        "logistics_status": "Delivered on 9/21",
    })
    assert [row["tracking_number"] for row in rows] == ["T1", "T2"]
    assert all(row["carrier"] is None and row["ship_date"] is None
               and row["delivery_status"] is None for row in rows)


class RecordingEmbedder(EmbeddingAdapter):
    def __init__(self, provider):
        super().__init__(Settings(environment="test", embedding_model="m", embedding_api_key="k",
                                  embedding_base_url="https://embed.invalid",
                                  embedding_provider=provider))
        self.contents = []

    def embed_contents(self, contents):
        self.contents.extend(contents)
        return [list(UNIT_VECTOR) for _ in contents]


def build_workbook() -> bytes:
    workbook = openpyxl.Workbook()
    orders = workbook.active
    orders.title = "订单信息"
    orders.append(["No.", "订单号", "SKU", "Product Description", "QTY", "购买日期", "质保结束日期",
                   "订单地址", "订单状态", "物流承运商", "物流单号", "物流状态", "发货日期", "订单金额"])
    orders.append([1, "Anker001", "E235311B", "eufy E25 combo", 1, datetime(2026, 9, 17),
                   datetime(2027, 9, 17), "addr", "已发货(拆成2个包裹发货）", "Fedex Amazon Logistic",
                   "Tracking number1: T1 Tracking number2: T1-1",
                   "Delivered on 9/21 Delivered on 9/23",
                   "2026-09-17 2026-09-19", 1099.99])
    orders.append([2, "Anker002", "T2080GA1", "S1 Pro Omni", 1, datetime(2026, 9, 13),
                   datetime(2026, 10, 13), "addr2", "已发货", "ANPOST", "T2", "Delivered",
                   datetime(2026, 9, 18), 1099.9])
    dealers = workbook.create_sheet("经销商信息")
    dealers.append(["country", "seller"])
    for country, seller in [("US", "best buy"), ("US", "Walmart"), ("DE", "brodos"),
                            ("DE", "Brodos AG"), ("DE", "Media Markt"), ("DE", "mediamarkt")]:
        dealers.append([country, seller])
    qa = workbook.create_sheet("QA")
    qa.append(["First degree directory", "Second degree directory", "Third degree directory",
               "Title", "Content", "Type", "Status", "Consumer Visible", "PN", "PDT"])
    qa.append(["Clean", "x", "E25/E28", "map lost", "The robot keeps losing the map.",
               "QA", "Published", "Visible", "T2352", "pdt"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def structured_env(database_engine, tmp_path):
    def purge(session):
        for document in session.scalars(select(KnowledgeDocument)).all():
            session.delete(document)
        session.commit()

    with Session(database_engine, expire_on_commit=False) as session:
        purge(session)
        seed_local_accounts(session)
        tenant = session.scalar(select(Tenant).where(Tenant.slug == "local-demo-a"))
        staff = session.scalar(select(StaffUser).where(StaffUser.external_subject == "local:staff"))
        session.commit()
        yield SimpleNamespace(engine=database_engine, storage=LocalEvidenceStorage(tmp_path),
                              tmp=tmp_path, tenant_id=tenant.id, staff_id=staff.id)
    with Session(database_engine) as session:
        purge(session)
        from serviceos_api.knowledge_ingest import ingest

        ingest(session)
        session.commit()


def upload_and_process(env, session, payload, filename, embedder, overrides=None):
    service = KnowledgePipelineService(
        session, Principal(env.tenant_id, env.staff_id, "STAFF", "客服", "SUPPORT"), env.storage)
    read, _ = service.upload(filename, payload, None, sheet_overrides=overrides)
    session.commit()
    job = session.scalar(select(KnowledgeJob).where(KnowledgeJob.document_id == read.document_id))
    error = process_knowledge_job(session, job, Settings(environment="test",
                                  evidence_storage_dir=str(env.tmp)), env.storage, embedder=embedder)
    session.commit()
    return read, error


def test_workbook_routing_facets_and_semantic_recall(structured_env):
    env = structured_env
    embedder = RecordingEmbedder("openai_compatible")
    with Session(env.engine, expire_on_commit=False) as session:
        read, error = upload_and_process(env, session, build_workbook(), "wb.xlsx", embedder)
        assert error is None
        segments = session.scalars(select(KnowledgeSegment).where(
            KnowledgeSegment.document_id == read.document_id)).all()
        orders = [s for s in segments if (s.attrs or {}).get("order_no")]
        dealers = [s for s in segments if (s.attrs or {}).get("seller_norm")]
        qa_rows = [s for s in segments if s.attrs == {}]
        assert len(orders) == 2 and len(dealers) == 6 and len(qa_rows) == 1
        first = next(s for s in orders if s.attrs["order_no"] == "anker001")
        assert first.attrs["warranty_end"] == "2027-09-17"
        assert first.attrs["carrier"] == "Fedex Amazon Logistic"
        assert first.attrs["shipments"] == [
            {"source_sequence": 1, "tracking_number": "T1", "carrier": None,
             "ship_date": "2026-09-17", "delivery_status": "Delivered on 9/21"},
            {"source_sequence": 2, "tracking_number": "T1-1", "carrier": None,
             "ship_date": "2026-09-19", "delivery_status": "Delivered on 9/23"},
        ]
        assert session.get(KnowledgeDocument, read.document_id).status == "EMBEDDED"
        retrieval = KnowledgeRetrieval(session)
        hits = retrieval.search("订单 承运商 Fedex 已发货", set())
        assert any((item.entry_key == first.entry_key) for item in hits.items)


def test_lookup_order_hits_warranty_boundary_and_misses_honestly(structured_env):
    env = structured_env
    embedder = RecordingEmbedder("openai_compatible")
    with Session(env.engine, expire_on_commit=False) as session:
        upload_and_process(env, session, build_workbook(), "wb.xlsx", embedder)
        tenant, conversation, run = active_run(session)
        service = AgentToolService(session, tenant.id, conversation.id, run.id, run.lease_token,
                                     settings=Settings(environment="test"))
        current = service.query_order_warranty("Anker001")
        assert current.status == "OK"
        current_payload = current.to_model_value()
        assert current_payload["warranty_end"] == "2027-09-17"
        assert current_payload["in_warranty"] == "yes"
        assert current_payload["data_scope"] == "REDACTED_STATIC_REFERENCE_NOT_LIVE_ORDER_DATA"
        assert current_payload["shipments"] == [
            {"source_sequence": 1, "tracking_number": "T1", "carrier": None,
             "ship_date": "2026-09-17", "delivery_status": "Delivered on 9/21"},
            {"source_sequence": 2, "tracking_number": "T1-1", "carrier": None,
             "ship_date": "2026-09-19", "delivery_status": "Delivered on 9/23"},
        ]
        expired = service.query_order_warranty("anker002")
        assert expired.to_model_value()["in_warranty"] == "no"
        missing = service.query_order_warranty("Anker999")
        assert missing.status == "NOT_FOUND"
        assert "参考数据中没有该订单号" in missing.to_prompt_text()


def test_match_dealer_three_states(structured_env):
    env = structured_env
    embedder = RecordingEmbedder("openai_compatible")
    with Session(env.engine, expire_on_commit=False) as session:
        upload_and_process(env, session, build_workbook(), "wb.xlsx", embedder)
        tenant, conversation, run = active_run(session)
        service = AgentToolService(session, tenant.id, conversation.id, run.id, run.lease_token,
                                     settings=Settings(environment="test"))
        assert service.check_dealer_authorization("US", "Best Buy").status == "OK"
        assert "AUTHORIZED" in service.check_dealer_authorization("德国", "Brodos AG").to_prompt_text()
        assert "AUTHORIZED" in service.check_dealer_authorization("DE", "brodos gmbh").to_prompt_text()
        assert service.check_dealer_authorization("DE", "mediamarkt").status == "DEGRADED"
        assert service.check_dealer_authorization("US", "unknown shop zz").status == "NOT_FOUND"
        assert service.check_dealer_authorization("JP", "best buy").status == "NOT_FOUND"


def test_image_document_uses_multimodal_content_and_openai_fallback(structured_env):
    env = structured_env
    with Session(env.engine, expire_on_commit=False) as session:
        multimodal = RecordingEmbedder("dashscope_multimodal")
        read, error = upload_and_process(env, session, PNG_BYTES, "label.png", multimodal)
        assert error is None
        segment = session.scalar(select(KnowledgeSegment).where(
            KnowledgeSegment.document_id == read.document_id))
        assert segment.image_key is not None
        assert session.get(KnowledgeDocument, read.document_id).status == "EMBEDDED"
        assert multimodal.contents and multimodal.contents[0]["image"].startswith("data:image/png;base64,")

        text_only = RecordingEmbedder("openai_compatible")
        read2, error2 = upload_and_process(env, session, PNG_BYTES + b"\x00", "label2.png", text_only)
        assert error2 is None
        segment2 = session.scalar(select(KnowledgeSegment).where(
            KnowledgeSegment.document_id == read2.document_id))
        assert session.get(KnowledgeDocument, read2.document_id).status == "STORED"
        assert session.scalar(select(KnowledgeEmbedding).where(
            KnowledgeEmbedding.segment_id == segment2.id)) is None
        assert text_only.contents == []


def test_sheet_overrides_force_generic_row(structured_env):
    env = structured_env
    embedder = RecordingEmbedder("openai_compatible")
    with Session(env.engine, expire_on_commit=False) as session:
        read, error = upload_and_process(env, session, build_workbook(), "wb.xlsx", embedder,
                                         overrides={"QA": "ROW"})
        assert error is None
        qa_segment = session.scalar(select(KnowledgeSegment).where(
            KnowledgeSegment.document_id == read.document_id,
            KnowledgeSegment.content.like("%Title: map lost%")))
        assert qa_segment is not None and qa_segment.attrs == {}


def test_all_agent_tools_share_one_envelope_schema(structured_env):
    env = structured_env
    embedder = RecordingEmbedder("openai_compatible")
    with Session(env.engine, expire_on_commit=False) as session:
        upload_and_process(env, session, build_workbook(), "wb.xlsx", embedder)
        tenant, conversation, run = active_run(session)
        service = AgentToolService(session, tenant.id, conversation.id, run.id, run.lease_token,
                                     settings=Settings(environment="test"))
        envelopes = [
            service.lookup_manual("map lost", verified_models=set()),
            service.query_order_warranty("Anker001"),
            service.query_order_warranty("nope-1"),
            service.check_dealer_authorization("US", "Walmart"),
            service.case_stage(),
            service.record_proposal("MISSING_EVIDENCE", "需要照片"),
            service.load_skill("diagnose-no-charge", "1.0.0"),
        ]
        for envelope in envelopes:
            assert envelope.status in TOOL_STATUSES
            assert set(envelope.audit("probe")) == {"tool", "status", "applicability", "references", "notes"}
            assert isinstance(envelope.items, list) and isinstance(envelope.notes, list)
        assert [e.status for e in envelopes] == ["OK", "OK", "NOT_FOUND", "OK", "NOT_FOUND",
                                                 "RECORDED", "LOADED"]
        db_session_refresh = session.get(AgentRun, run.id)
        tools = [entry["tool"] for entry in db_session_refresh.result_payload["tool_envelopes"]]
        assert tools == ["lookup_manual", "query_order_warranty", "query_order_warranty",
                         "check_dealer_authorization", "get_case_stage", "record_proposal", "load_skill"]


def test_shipping_intent_replay_gets_warranty_and_dealer_conclusions(structured_env):
    """Replay of hackathon shipping-inquiry complaints (客诉工单信息 rows 2-4: Anker001-003)."""
    env = structured_env
    with Session(env.engine, expire_on_commit=False) as session:
        from serviceos_api.knowledge_ingest import ingest
        ingest(session)
        session.commit()
        tenant, conversation, run = active_run(session)
        service = AgentToolService(session, tenant.id, conversation.id, run.id, run.lease_token,
                                     settings=Settings(environment="test"))
        for order_no, expected in (("Anker001", "yes"), ("Anker002", "yes"), ("Anker003", "yes")):
            envelope = service.query_order_warranty(order_no)
            assert envelope.status == "OK"
            assert envelope.to_model_value()["in_warranty"] == expected
        split_order = service.query_order_warranty("Anker001").to_model_value()
        assert [shipment["tracking_number"] for shipment in split_order["shipments"]] == [
            "Anker1234567", "Anker1234567-1",
        ]
        assert [shipment["delivery_status"] for shipment in split_order["shipments"]] == [
            "Delivered on 9/21", "Delivered on 9/23",
        ]
        assert split_order["carrier_source_text"] == "Fedex Amazon Logistic"
        dealer = service.check_dealer_authorization("US", "best buy")
        assert dealer.status == "OK" and "AUTHORIZED" in dealer.to_prompt_text()


def test_staff_search_and_stats_reuse_single_retrieval(structured_env):
    env = structured_env
    with Session(env.engine, expire_on_commit=False) as session:
        from serviceos_api.knowledge_ingest import ingest
        ingest(session)
        session.commit()
        service = KnowledgePipelineService(
            session, Principal(env.tenant_id, env.staff_id, "STAFF", "客服", "SUPPORT"), env.storage)
        hits = service.search("robot keeps losing the map", 5)
        assert hits and hits[0].entry_key.startswith("qa-v1:robovac:")
        stats = service.stats()
        assert stats.corpus_present and stats.segments >= 1235 and stats.documents >= 6
